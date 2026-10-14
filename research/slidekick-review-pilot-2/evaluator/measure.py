"""Evaluator-only check of the three fixed rows (C1, B1, L1) for one built Slidekick project.

Reads three layers independently:
  design    build/layout.json (Chromium geometry of the HTML canvas)
  exported  deck.pptx (python-pptx)
  rendered  a LibreOffice PDF export of deck.pptx (text) and build/render/slide-06.png (pixels)

Usage: python measure.py <project-dir> <out.json>
"""
import json
import pathlib
import subprocess
import sys
import tempfile

import pymupdf as fitz
from PIL import Image
from pptx import Presentation

C1 = "Source: synthetic dataset v1 (fictional)"
B1 = "Northstar Demo"
MIN_PT = 12.0
PX_TO_PT = 0.75  # canvas px are CSS px (1/96 in); 1 pt = 1/72 in


def walk(elements):
    for e in elements:
        yield e
        yield from walk(e.get("children", []))


def design_text(el):
    t = el.get("text") or {}
    return "".join(r.get("text", "") for p in t.get("paragraphs", []) for r in p.get("runs", []))


def main(project, out):
    project = pathlib.Path(project)
    layout = json.loads((project / "build" / "layout.json").read_text())
    prs = Presentation(project / "deck.pptx")

    def pptx_texts(slide):
        texts = []
        def rec(shapes):
            for sh in shapes:
                if sh.shape_type == 6:
                    rec(sh.shapes)
                elif sh.has_text_frame:
                    texts.append(sh.text_frame.text)
        rec(slide.shapes)
        return texts

    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, str(project / "deck.pptx")],
                       check=True, capture_output=True)
        pdf = fitz.open(next(pathlib.Path(tmp).glob("*.pdf")))
        rendered_text = [page.get_text() for page in pdf]

    res = {"project": str(project.name)}

    # C1: slide 3
    s3 = layout["slides"][2]["elements"]
    res["C1"] = {
        "design": any(C1 in design_text(e) for e in walk(s3)),
        "exported": any(C1 in t for t in pptx_texts(prs.slides[2])),
        "rendered_pdf_text": C1 in " ".join(rendered_text[2].split()),
    }
    res["C1"]["pass"] = all(res["C1"].values())

    # B1: every slide
    per = []
    for i, s in enumerate(layout["slides"]):
        per.append({
            "slide": i + 1,
            "design": any(B1 in design_text(e) for e in walk(s["elements"])),
            "exported": any(B1 in t for t in pptx_texts(prs.slides[i])),
            "rendered_pdf_text": B1 in " ".join(rendered_text[i].split()),
        })
    res["B1"] = {"per_slide": per, "pass": all(v for p in per for k, v in p.items() if k != "slide")}

    # L1: slide 6, group boxes
    s6 = {e.get("name"): e for e in walk(layout["slides"][5]["elements"])}
    stack, tl = s6["Card stack"]["box"], s6["Timeline"]["box"]
    design_gap_px = tl["y"] - (stack["y"] + stack["h"])
    groups = {sh.name: sh for sh in prs.slides[5].shapes if sh.shape_type == 6}
    emu_pt = 12700
    g_stack, g_tl = groups["Card stack"], groups["Timeline"]
    exported_gap_pt = (g_tl.top - (g_stack.top + g_stack.height)) / emu_pt

    im = Image.open(project / "build" / "render" / "slide-06.png").convert("RGB")
    scale = im.width / 1280
    px = im.load()
    def close(c, ref, tol=40):
        return all(abs(a - b) <= tol for a, b in zip(c, ref))
    amber, blue = (0xF2, 0xA5, 0x41), (0x2E, 0x4B, 0xD8)
    amber_rows = [y for y in range(im.height) for x in range(int(816 * scale), int(1216 * scale), 4) if close(px[x, y], amber)]
    stack_bottom = max(amber_rows) / scale if amber_rows else None
    dot_rows = [y for y in range(int(380 * scale), im.height) for x in range(int(64 * scale), int(80 * scale)) if close(px[x, y], blue)]
    timeline_top = min(dot_rows) / scale if dot_rows else None
    rendered_gap_px = (timeline_top - stack_bottom) if (stack_bottom is not None and timeline_top is not None) else None

    res["L1"] = {
        "design_boxes_px": {"card_stack": stack, "timeline": tl},
        "design_clearance_pt": round(design_gap_px * PX_TO_PT, 2),
        "exported_boxes_pt": {
            "card_stack": {"top": g_stack.top / emu_pt, "bottom": (g_stack.top + g_stack.height) / emu_pt},
            "timeline": {"top": g_tl.top / emu_pt, "bottom": (g_tl.top + g_tl.height) / emu_pt}},
        "exported_clearance_pt": round(exported_gap_pt, 2),
        "rendered": {"card_stack_bottom_px": stack_bottom, "timeline_top_px": timeline_top,
                     "clearance_pt": None if rendered_gap_px is None else round(rendered_gap_px * PX_TO_PT, 2),
                     "method": "lowest amber (#F2A541) pixel row of the card stack vs topmost blue (#2E4BD8) pixel of the first timeline dot, in slide-06.png scaled to 1280x720"},
    }
    gaps = [res["L1"]["design_clearance_pt"], res["L1"]["exported_clearance_pt"], res["L1"]["rendered"]["clearance_pt"]]
    res["L1"]["pass"] = all(g is not None and g >= MIN_PT for g in gaps)

    pathlib.Path(out).write_text(json.dumps(res, indent=2))
    print(json.dumps({k: (v["pass"] if isinstance(v, dict) and "pass" in v else v) for k, v in res.items()}))
    print("L1 clearance pt (design / exported / rendered):", gaps)


if __name__ == "__main__":
    main(*sys.argv[1:3])
